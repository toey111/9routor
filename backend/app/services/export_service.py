import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def export_personnel_csv(personnel_list: list) -> io.StringIO:
    output = io.StringIO()
    # Write UTF-8 BOM for Thai language display in Excel
    output.write('\ufeff')
    writer = csv.writer(output)
    
    headers = [
        "รหัสบุคลากร", "คำนำหน้า", "ชื่อ", "นามสกุล", "ตำแหน่งงาน",
        "สาขาวิชา/หน่วยงาน", "ประเภทบุคลากร", "สถานะการทำงาน", "อีเมล", "เบอร์โทรศัพท์"
    ]
    writer.writerow(headers)
    
    for p in personnel_list:
        writer.writerow([
            p.personnel_code,
            p.prefix_th,
            p.first_name_th,
            p.last_name_th,
            p.position.name if p.position else "-",
            p.branch.name if p.branch else (p.department.name if p.department else "-"),
            p.personnel_type.name if p.personnel_type else "-",
            p.work_status or "-",
            p.email,
            p.phone or "-"
        ])
    output.seek(0)
    return output

def export_personnel_excel(personnel_list: list) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "รายชื่อบุคลากร"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers = [
        "ลำดับ", "รหัสบุคลากร", "ชื่อ-นามสกุล", "ตำแหน่ง",
        "สาขาวิชา", "ประเภทบุคลากร", "สถานะ", "อีเมล", "เบอร์โทรศัพท์", "วันที่บรรจุ"
    ]
    
    ws.append(headers)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = thin_border

    for idx, p in enumerate(personnel_list, start=1):
        full_name = f"{p.prefix_th}{p.first_name_th} {p.last_name_th}"
        row = [
            idx,
            p.personnel_code,
            full_name,
            p.position.name if p.position else "-",
            p.branch.name if p.branch else "-",
            p.personnel_type.name if p.personnel_type else "-",
            p.work_status or "-",
            p.email,
            p.phone or "-",
            str(p.appoint_date) if p.appoint_date else "-"
        ]
        ws.append(row)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=idx + 1, column=col_num)
            cell.border = thin_border

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_personnel_pdf(personnel_list: list) -> io.BytesIO:
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )
    
    elements = []
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=14,
        alignment=1, # Center
        spaceAfter=15
    )
    
    title = Paragraph("รายงานข้อมูลบุคลากร คณะเทคโนโลยีสารสนเทศ มหาวิทยาลัยราชภัฏร้อยเอ็ด", title_style)
    elements.append(title)
    elements.append(Spacer(1, 10))
    
    data = [
        ["#", "รหัส", "ชื่อ - นามสกุล", "ตำแหน่ง", "สาขาวิชา", "ประเภท", "สถานะ", "อีเมล"]
    ]
    
    for idx, p in enumerate(personnel_list, start=1):
        full_name = f"{p.prefix_th}{p.first_name_th} {p.last_name_th}"
        data.append([
            str(idx),
            p.personnel_code,
            full_name,
            p.position.name if p.position else "-",
            p.branch.name if p.branch else "-",
            p.personnel_type.name if p.personnel_type else "-",
            p.work_status or "-",
            p.email
        ])
        
    t = Table(data, colWidths=[30, 60, 160, 110, 140, 110, 70, 140])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (2, 1), (2, -1), 'LEFT'),
        ('ALIGN', (7, 1), (7, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    
    elements.append(t)
    doc.build(elements)
    output.seek(0)
    return output
